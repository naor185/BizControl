"""
CSV / Excel connector — for any system that can export a file, including ones with no API.
The business owner uploads the file and maps its columns; everything after that is the engine.
"""
from __future__ import annotations

import csv
import io
from datetime import datetime

from app.connectors.base import BaseConnector, ConnectorError, ConnectorManifest, SourceTable

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 50_000
MAX_COLUMNS = 150


class CsvExcelConnector(BaseConnector):
    manifest = ConnectorManifest(
        slug="csv_excel",
        name="קובץ Excel / CSV",
        icon="file-spreadsheet",
        version="1.0",
        auth_type="file",
        supported_entities=("clients", "services"),
        description="ייצוא מכל מערכת אחרת (Arbox, קלמארק, Excel ידני…) לקובץ, והעלאה כאן",
        requires_mapping=True,
        unsupported=("payment_token_migration",),
    )

    def __init__(self, filename: str, content: bytes, sheet: str | None = None):
        self.filename = filename or ""
        self.content = content
        self.sheet = sheet

    def read(self, entity: str) -> SourceTable:
        if len(self.content) > MAX_FILE_BYTES:
            raise ConnectorError("הקובץ גדול מ-10MB — פצל אותו לכמה קבצים")
        if not self.content:
            raise ConnectorError("הקובץ ריק")
        name = self.filename.lower()
        if name.endswith(".xlsx") or name.endswith(".xlsm") or self.content[:2] == b"PK":
            headers, rows, info = self._read_xlsx()
        elif name.endswith(".xls"):
            raise ConnectorError("קובץ Excel בפורמט הישן (.xls) לא נתמך — פתח אותו ושמור כ-xlsx או כ-CSV")
        else:
            headers, rows, info = self._read_csv()
        if not headers:
            raise ConnectorError("לא נמצאה שורת כותרות בקובץ")
        if len(headers) > MAX_COLUMNS:
            raise ConnectorError(f"יש בקובץ {len(headers)} עמודות — המקסימום הוא {MAX_COLUMNS}")
        if not rows:
            raise ConnectorError("אין בקובץ שורות מתחת לשורת הכותרות")
        if len(rows) > MAX_ROWS:
            raise ConnectorError(f"יש בקובץ יותר מ-{MAX_ROWS:,} שורות — המקסימום לייבוא אחד. פצל לכמה קבצים")
        return SourceTable(entity=entity, headers=headers, rows=rows, info=info)

    # ── CSV ──────────────────────────────────────────────────────────────────
    def _read_csv(self) -> tuple[list[str], list[list], dict]:
        text, encoding = None, None
        for enc in ("utf-8-sig", "cp1255"):   # cp1255: what Hebrew Excel saves "CSV" as
            try:
                text, encoding = self.content.decode(enc), enc
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ConnectorError("לא הצלחנו לקרוא את הקידוד של הקובץ — שמור אותו מחדש כ-CSV UTF-8")
        delimiter = _detect_delimiter(text)
        table = [r for r in csv.reader(io.StringIO(text), delimiter=delimiter)]
        headers, rows = _split_header(table)
        return headers, rows, {"encoding": encoding, "delimiter": {"\t": "tab"}.get(delimiter, delimiter)}

    # ── Excel ────────────────────────────────────────────────────────────────
    def _read_xlsx(self) -> tuple[list[str], list[list], dict]:
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(self.content), read_only=True, data_only=True)
        except Exception:
            raise ConnectorError("לא הצלחנו לפתוח את קובץ ה-Excel — ודא שהוא לא פגום ולא מוגן בסיסמה")
        try:
            names = wb.sheetnames
            chosen = self.sheet if self.sheet in names else None
            table: list[list] = []
            for sheet_name in ([chosen] if chosen else names):
                ws = wb[sheet_name]
                table = []
                for i, r in enumerate(ws.iter_rows(values_only=True)):
                    if i > MAX_ROWS + 50:
                        break
                    table.append([_xlsx_value(v) for v in r])
                if any(any(v not in (None, "") for v in r) for r in table):
                    chosen = sheet_name
                    break
            headers, rows = _split_header(table)
            return headers, rows, {"sheet": chosen, "sheets": names}
        finally:
            wb.close()


def _detect_delimiter(text: str) -> str:
    """The separator that appears the same number of times on most of the first lines (the header
    line must have it). csv.Sniffer misreads Hebrew files with ';', which Excel uses in many locales."""
    lines = [ln for ln in text.splitlines()[:30] if ln.strip()]
    best, best_score = ",", -1
    for d in (",", ";", "\t", "|"):
        counts = [ln.count(d) for ln in lines]
        if not counts or counts[0] == 0:
            continue
        score = sum(1 for c in counts if c == counts[0]) * 1000 + counts[0]
        if score > best_score:
            best, best_score = d, score
    return best


def _xlsx_value(v):
    # Keep dates as dates (unambiguous), everything else as the cell shows it.
    if isinstance(v, datetime) and v.hour == 0 and v.minute == 0 and v.second == 0:
        return v.date()
    return v


def _split_header(table: list[list]) -> tuple[list[str], list[list]]:
    """First non-empty row = headers; empty rows dropped; blank/duplicate headers get a name."""
    table = [r for r in table if any(v not in (None, "") and str(v).strip() for v in r)]
    if not table:
        return [], []
    raw_headers = table[0]
    width = max(len(r) for r in table)
    # drop trailing columns that are empty everywhere
    while width > 0 and all((len(r) < width or r[width - 1] in (None, "")) for r in table):
        width -= 1
    headers, seen = [], {}
    for i in range(width):
        h = str(raw_headers[i]).strip() if i < len(raw_headers) and raw_headers[i] not in (None, "") else f"עמודה {i + 1}"
        if h in seen:
            seen[h] += 1
            h = f"{h} ({seen[h]})"
        else:
            seen[h] = 1
        headers.append(h)
    rows = [[(r[i] if i < len(r) else None) for i in range(width)] for r in table[1:]]
    return headers, rows


CONNECTOR = CsvExcelConnector
